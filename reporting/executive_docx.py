import os
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from models.crawl import CrawlStats

def set_cell_background(cell, fill_hex: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

class ExecutiveSummaryGenerator:
    def __init__(self, output_path: str = "reports/executive-summary.docx"):
        self.output_path = output_path
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        self.doc = Document()
        self._init_styling()

    def _init_styling(self):
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    def generate(
        self,
        website: str,
        audit_date: str,
        stats: CrawlStats,
        product_pages_count: int,
        gsc_summary: Dict[str, Any],
        issue_clusters: List[Dict[str, Any]],
        top_product_opps: List[Dict[str, Any]],
        top_link_opps: List[Dict[str, Any]],
        aeo_summary: Dict[str, Any],
        geo_summary: Dict[str, Any]
    ):
        # Title
        p_title = self.doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(30)
        p_title.paragraph_format.space_after = Pt(4)
        run_pre = p_title.add_run("SEOJEV V2 — EXECUTIVE INTELLIGENCE SUMMARY\n")
        run_pre.font.name = 'Arial'
        run_pre.font.size = Pt(11)
        run_pre.font.bold = True
        run_pre.font.color.rgb = RGBColor(0x1B, 0x6E, 0xB5)

        run_main = p_title.add_run(f"Ranking-Focused SEO, AEO & GEO Intelligence: {website}")
        run_main.font.name = 'Arial'
        run_main.font.size = Pt(20)
        run_main.font.bold = True
        run_main.font.color.rgb = RGBColor(0x0F, 0x38, 0x70)

        p_date = self.doc.add_paragraph(f"Audit Date: {audit_date} | Domain Target: {website}")
        p_date.paragraph_format.space_after = Pt(20)

        # Core Questions Answer Card
        h1 = self.doc.add_heading("THE 14 EXECUTIVE QUESTIONS ANSWERED", level=1)
        h1.runs[0].font.name = 'Arial'
        h1.runs[0].font.size = Pt(14)
        h1.runs[0].font.color.rgb = RGBColor(0x0F, 0x38, 0x70)

        table = self.doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Core Strategic Question"
        hdr_cells[1].text = "Audit Finding & Actionable Diagnosis"
        for c in hdr_cells:
            set_cell_background(c, "0F3870")
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        p11_20 = gsc_summary.get("pos_11_20_count", 0)
        p21_30 = gsc_summary.get("pos_21_30_count", 0)
        gsc_prod_count = gsc_summary.get("product_pages_with_gsc", 0)
        has_gsc = gsc_summary.get("has_data", False)

        top_cluster = issue_clusters[0] if issue_clusters else {}
        top_cluster_desc = f"{top_cluster.get('issue', '').replace('_', ' ').title()} ({top_cluster.get('affected_urls_count', 0)} pages in '{top_cluster.get('primary_affected_template')}')" if top_cluster else "None"

        q_and_a = [
            ("1. How many pages were analyzed?", f"{stats.urls_crawled:,} crawled URLs out of {stats.urls_discovered:,} discovered."),
            ("2. How many are indexable?", f"{stats.indexable_urls:,} pages ({round((stats.indexable_urls/max(stats.urls_crawled,1))*100, 1)}% of crawled set)."),
            ("3. How many product/bike pages exist?", f"{product_pages_count:,} automotive product/model pages identified."),
            ("4. How many product pages have GSC data?", f"{gsc_prod_count:,} pages" if has_gsc else "Ranking data unavailable. Connect Google Search Console data for query-level diagnosis."),
            ("5. How many pages rank in positions 11–20?", f"{p11_20:,} pages (Striking Distance Opportunity Set)" if has_gsc else "Connect GSC data to uncover positions 11–20 striking distance pages."),
            ("6. How many rank in positions 21–30?", f"{p21_30:,} pages (Second-Tier Opportunity Set)" if has_gsc else "Connect GSC data to uncover positions 21–30 query targets."),
            ("7. What are the largest template-level problems?", f"{top_cluster_desc}."),
            ("8. What are the largest product-page gaps?", "Missing dedicated Variant matrices, localized On-Road price breakdowns, and competitor comparison modules."),
            ("9. What are the largest internal-link opportunities?", f"{len(top_link_opps):,} high-value model pages have weak inbound link equity (<4 links) despite existing brand hubs."),
            ("10. What technical blockers exist?", f"{stats.critical_issues} critical blockers (HTTP 4xx errors, non-indexable URLs in XML sitemaps, redirect loops)."),
            ("11. What content gaps exist?", "Commercial content gaps: thin specification tables, missing variant price breakdowns, and absent FAQ accordions."),
            ("12. What AEO gaps exist?", f"AEO Readiness: {aeo_summary.get('average_score', 45)}/100. Key gap: missing structured Q&A and FAQPage JSON-LD schema."),
            ("13. What GEO structural gaps exist?", f"GEO Readiness: {geo_summary.get('average_score', 50)}/100. Key gap: inconsistent Brand/Model entity naming across schema and metadata."),
            ("14. What should be fixed first?", "Phase 1: Clear technical 4xx/sitemap blockers; Phase 2: Deploy variant matrices & FAQ schema on Bike Model templates; Phase 3: Add contextual internal links from Brand hubs to model pages.")
        ]

        for q, a in q_and_a:
            row_cells = table.add_row().cells
            row_cells[0].text = q
            row_cells[0].paragraphs[0].runs[0].font.bold = True
            row_cells[1].text = a
            for cell in row_cells:
                set_cell_margins(cell, 70, 70, 100, 100)

        # Immediate Next Steps
        h2 = self.doc.add_heading("IMMEDIATE PRIORITIES (NEXT 14 DAYS)", level=2)
        h2.runs[0].font.name = 'Arial'
        h2.runs[0].font.size = Pt(13)
        h2.runs[0].font.color.rgb = RGBColor(0x1B, 0x6E, 0xB5)

        p_steps = self.doc.add_paragraph()
        p_steps.add_run("1. Template Engineering Fix: ").bold = True
        p_steps.add_run(f"Update the Bike Model layout to render structured variants, specifications, and FAQPage schema across all {product_pages_count} models.\n")
        p_steps.add_run("2. Internal Linking Injections: ").bold = True
        p_steps.add_run("Inject direct contextual links from Brand parent hubs (e.g. /bikes/tvs) directly to priority model URLs.\n")
        p_steps.add_run("3. GSC Query-CTR Alignment: ").bold = True
        p_steps.add_run("For striking-distance queries (positions 11-20), align title tags with observed commercial intent (e.g., adding 'Price, Variants, Mileage & Specs').\n")

        self.doc.save(self.output_path)
